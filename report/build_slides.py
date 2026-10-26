"""
report/build_slides.py
----------------------
Builds the end-semester review deck (report/B9_Final_Review.pptx) with python-pptx.

Inputs: report/figures/{doc,slides}/*.png and report/numbers.json (run report/make_figures.py first).
Usage:  python report/build_slides.py
"""

import os
import json
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DOC = os.path.join(HERE, 'figures', 'doc')
FIG_SLIDE = os.path.join(HERE, 'figures', 'slides')
OUT = os.path.join(HERE, 'B9_Final_Review.pptx')
N = json.load(open(os.path.join(HERE, 'numbers.json')))

TITLE = 'Semantic Deep JSCC over 6G Fading Channels'
TEAM = 'Team B-9'
MEMBERS = ['Rithvik Arulprakash', 'Harshith K V', 'Venugopalan G', 'Vipin Sudhakar']

# Palette: deep navy for title/closing slides, white content slides; blue = our system,
# orange/aqua = digital baselines (same as the figures).
NAVY = RGBColor(0x10, 0x23, 0x3D)
NAVY_2 = RGBColor(0x1B, 0x35, 0x58)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
INK = RGBColor(0x1A, 0x1F, 0x2B)
INK_2 = RGBColor(0x4B, 0x55, 0x63)
MUTED = RGBColor(0x6B, 0x72, 0x80)
CARD = RGBColor(0xF1, 0xF5, 0xFA)
CARD_EDGE = RGBColor(0xD9, 0xE2, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xC9, 0xDB, 0xF2)
FONT = 'Calibri'

SW, SH = Inches(13.333), Inches(7.5)
MARGIN = 0.6


def psnr(method, ch, snr):
    return N['psnr'][method][ch][str(snr)]


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_spacing=1.1, name=None):
    """runs: str, or list of paragraphs; each paragraph is a str or a list of (text, {opts})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for t, opts in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get('size', size))
            f.bold = opts.get('bold', bold)
            f.italic = opts.get('italic', False)
            f.color.rgb = opts.get('color', color)
            if opts.get('sub'):
                r._r.get_or_add_rPr().set('baseline', '-25000')
            if opts.get('sup'):
                r._r.get_or_add_rPr().set('baseline', '30000')
    return box


def bullets(slide, x, y, w, h, items, size=18, color=INK, gap=8):
    """Real bulleted paragraphs (one per item). Items may be str or run lists."""
    box = text(slide, x, y, w, h, items, size=size, color=color)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Emu(Inches(0.28))))
        pPr.set('indent', str(-Emu(Inches(0.28))))
        for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
            for el in pPr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1]):
                pPr.remove(el)
        from pptx.oxml.ns import qn
        bu_clr = pPr.makeelement(qn('a:buClr'), {})
        clr = bu_clr.makeelement(qn('a:srgbClr'), {'val': '2A78D6'})
        bu_clr.append(clr)
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•'}))
        p.space_after = Pt(gap)
    return box


def rect(slide, x, y, w, h, fill=CARD, line=None, radius=True, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def circle_num(slide, x, y, n, d=0.55, fill=BLUE, color=WHITE, size=18):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def arrow(slide, x, y, w=0.45, h=0.4, color=MUTED):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def picture(slide, path, x, y, max_w, max_h, align='center'):
    """Insert an image scaled to fit the box, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    px = x + (max_w - w) / 2 if align == 'center' else x
    py = y + (max_h - h) / 2
    return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(w), Inches(h))


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


# ------------------------------------------------------------------ slide frames
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    def slide(self, notes, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        self.count += 1
        background(s, NAVY if dark else WHITE)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def content(self, kicker, title, notes):
        s = self.slide(notes)
        text(s, MARGIN, 0.38, 12.1, 0.3, kicker.upper(), size=12, color=BLUE, bold=True, name='Kicker')
        text(s, MARGIN, 0.68, 12.1, 0.8, title, size=32, color=INK, bold=True, name='Title')
        text(s, MARGIN, 7.02, 9, 0.3, f'{TEAM}  |  {TITLE}', size=10, color=MUTED, name='Footer')
        text(s, 12.733 - 1, 7.02, 1, 0.3, str(self.count), size=10, color=MUTED, align=PP_ALIGN.RIGHT,
             name='Slide number')
        return s

    def save(self):
        self.prs.save(OUT)


def card(slide, x, y, w, h, heading, body, heading_color=INK, size=16, fill=CARD, num=None):
    rect(slide, x, y, w, h, fill=fill)
    pad = 0.25
    hx = x + pad
    if num is not None:
        circle_num(slide, x + pad, y + pad, num, d=0.5, size=16)
        hx = x + pad + 0.65
    text(slide, hx, y + pad + (0.07 if num is not None else 0), w - (hx - x) - pad, 0.45, heading,
         size=19, color=heading_color, bold=True)
    body_y = y + pad + 0.62
    if isinstance(body, list) and body and not isinstance(body[0], tuple):
        bullets(slide, x + pad, body_y, w - 2 * pad, h - (body_y - y) - pad, body, size=size, gap=6)
    else:
        text(slide, x + pad, body_y, w - 2 * pad, h - (body_y - y) - pad, body, size=size, color=INK_2)


# ------------------------------------------------------------------ deck
def build():
    d = Deck()

    # 1. Title ----------------------------------------------------------------
    s = d.slide(
        'Introduce the team and the topic: transmitting camera frames over simulated 6G wireless channels '
        'with deep joint source-channel coding. One model handles three channel conditions: AWGN, Rician '
        '(line of sight) and Rayleigh (non line of sight).', dark=True)
    text(s, 0.9, 1.3, 11.5, 0.4, 'END-SEMESTER PROJECT REVIEW', size=14, color=ICE, bold=True)
    text(s, 0.9, 1.75, 11.5, 1.75, TITLE, size=48, color=WHITE, bold=True, line_spacing=1.0)
    text(s, 0.9, 3.65, 11.5, 0.6,
         'Real-time image transmission that stays robust under AWGN, Rician and Rayleigh fading',
         size=20, color=ICE)
    for i, ch in enumerate(['AWGN', 'Rician (LoS)', 'Rayleigh (NLoS)']):
        chip = rect(s, 0.9 + i * 2.25, 4.45, 2.05, 0.48, fill=NAVY_2)
        text(s, 0.9 + i * 2.25, 4.45, 2.05, 0.48, ch, size=14, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.9, 5.55, 11.5, 0.45, TEAM, size=20, color=WHITE, bold=True)
    text(s, 0.9, 6.0, 11.5, 0.45, '   ·   '.join(MEMBERS), size=16, color=ICE)
    text(s, 0.9, 6.5, 11.5, 0.4, date.today().strftime('%B %Y'), size=13, color=ICE)

    # 2. Outline --------------------------------------------------------------
    s = d.content('Overview', 'Outline',
                  'Walk through the structure: why the problem matters, what exists already, how our system '
                  'is built and trained, how the project evolved, the three main results, and conclusions.')
    items = [('Motivation and objectives', 'Why joint source-channel coding for 6G'),
             ('Background', 'Deep JSCC and related work'),
             ('System design', 'Pipeline, channel and receiver models'),
             ('Training and evolution', 'Methodology and how the design changed'),
             ('Results', 'Equalization, bandwidth, comparison with digital'),
             ('Conclusion', 'Contributions, limitations, future work')]
    for i, (h, b) in enumerate(items):
        col, row = i % 3, i // 3
        x, y = MARGIN + col * 4.1, 1.85 + row * 2.45
        rect(s, x, y, 3.85, 2.15)
        circle_num(s, x + 0.3, y + 0.3, i + 1)
        text(s, x + 0.3, y + 1.0, 3.3, 0.5, h, size=20, bold=True)
        text(s, x + 0.3, y + 1.48, 3.3, 0.6, b, size=15, color=INK_2)

    # 3. Motivation -----------------------------------------------------------
    s = d.content('Motivation', 'Why rethink how images are sent over wireless links',
                  'Conventional systems compress first and protect bits second. That is near-optimal at the '
                  'SNR they are designed for, but fails abruptly below it (the cliff effect) and does not '
                  'improve above it. Deep JSCC learns a direct mapping from pixels to channel symbols, so '
                  'quality degrades gradually. 6G links switch between line-of-sight and non-line-of-sight '
                  'conditions, so robustness across fading types matters.')
    for i, (head, color, body) in enumerate([
        ('Conventional digital transmission', ORANGE,
         ['Compress (JPEG / WebP), then add a channel code', 'Near-optimal at its design SNR',
          'Cliff effect: below the threshold, decoding fails completely', 'Above it, quality stops improving']),
        ('Deep joint source-channel coding', BLUE,
         ['A neural network maps pixels directly to channel symbols', 'Quality degrades gracefully as SNR falls',
          'Adapts to the channel it was trained on', 'A building block of semantic communication for 6G']),
    ]):
        x = MARGIN + i * 6.15
        rect(s, x, 1.75, 5.95, 3.85)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(2.08), Inches(0.24), Inches(0.24))
        dot.fill.solid(); dot.fill.fore_color.rgb = color; dot.line.fill.background(); dot.shadow.inherit = False
        text(s, x + 0.68, 1.98, 5.0, 0.45, head, size=20, bold=True)
        bullets(s, x + 0.3, 2.7, 5.4, 2.8, body, size=17)
    rect(s, MARGIN, 5.85, 12.1, 0.95, fill=RGBColor(0xE8, 0xF1, 0xFC))
    text(s, MARGIN + 0.3, 5.85, 11.5, 0.95,
         [[('6G challenge: ', {'bold': True, 'color': BLUE}),
           ('links switch between line-of-sight (Rician) and non-line-of-sight (Rayleigh) fading, so the '
            'transmission scheme must stay robust across channel types and signal qualities.', {})]],
         size=17, anchor=MSO_ANCHOR.MIDDLE)

    # 4. Objectives -----------------------------------------------------------
    s = d.content('Problem statement', 'Objectives',
                  'Four objectives: a real-time system for camera frames; one trained model for all three '
                  'channel types; graceful degradation instead of a cliff; and a quantitative study of the '
                  'bandwidth and SNR trade-offs, including a comparison with digital transmission.')
    objs = [('Real-time transmission', 'Send 128x128 camera frames through a simulated wireless channel '
                                       'and reconstruct them live'),
            ('One model, three channels', 'A single trained model for AWGN, Rician (LoS) and Rayleigh '
                                          '(NLoS) fading'),
            ('Graceful degradation', 'Image quality should fall smoothly as SNR drops, with no cliff'),
            ('Quantify trade-offs', 'Measure bandwidth vs quality and compare against digital transmission '
                                    'at the same bandwidth')]
    for i, (h, b) in enumerate(objs):
        col, row = i % 2, i // 2
        card(s, MARGIN + col * 6.15, 1.8 + row * 2.55, 5.95, 2.3, h, b, num=i + 1, size=17)

    # 5. Related work ---------------------------------------------------------
    s = d.content('Background', 'Related work',
                  'Deep JSCC was introduced by Bourtsoulatze, Kurka and Gunduz in 2019 with a CNN autoencoder. '
                  'ADJSCC added attention so one network covers a range of SNRs. WITT and SwinJSCC moved to '
                  'transformers. Recent diffusion-based methods produce realistic images but need many '
                  'denoising steps. Our system works in the latent space of a pretrained VAE, uses a ViT, '
                  'and adds explicit equalization so one model handles three channel types in real time.')
    rows = [('Work', 'Key idea', 'Limitation'),
            ('Deep JSCC — Bourtsoulatze et al., 2019', 'CNN autoencoder maps images to channel symbols',
             'Pixel domain, small images'),
            ('ADJSCC — Xu et al., 2022', 'Attention adapts one network to a range of SNRs', 'CNN, pixel domain'),
            ('WITT — Yang et al., 2023', 'Vision transformer for JSCC', 'Larger models, image only'),
            ('Diffusion JSCC (2024–25)', 'Generative decoders for realistic images', 'Many denoising steps: slow'),
            ('This project', 'Frozen VAE latent + ViT JSCC + MMSE-equalized receiver; real time',
             'Bounded by VAE quality')]
    tbl = s.shapes.add_table(len(rows), 3, Inches(MARGIN), Inches(1.75), Inches(12.1), Inches(4.9)).table
    for c, wdt in enumerate([4.1, 4.9, 3.1]):
        tbl.columns[c].width = Inches(wdt)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (RGBColor(0xE8, 0xF1, 0xFC) if r == len(rows) - 1
                                                            else (CARD if r % 2 else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.15)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ''
            run = p.add_run()
            run.text = val
            run.font.name, run.font.size = FONT, Pt(15)
            run.font.bold = r == 0 or r == len(rows) - 1 and c == 0
            run.font.color.rgb = WHITE if r == 0 else INK

    # 6. Architecture ---------------------------------------------------------
    s = d.content('System design', 'System architecture',
                  'Left to right: the camera frame is compressed by a frozen Stable Diffusion VAE into a 4x16x16 '
                  'latent. A vision transformer maps it to K complex channel symbols with unit average power. '
                  'The channel applies fading and noise. The receiver equalizes each symbol with MMSE using the '
                  'channel state, a second transformer reconstructs the latent, and the frozen VAE decodes the '
                  'image. Only the two transformers are trained.')
    picture(s, os.path.join(FIG_SLIDE, 'fig1_system.png'), MARGIN, 1.7, 12.1, 2.55)
    for i, (h, b, col) in enumerate([
        ('Frozen VAE', 'Compresses a 128x128 frame to a 4x16x16 latent: 48x fewer values. Not trained.', INK),
        ('ViT transmitter', '64 patch tokens -> K complex symbols, normalized to unit average power.', BLUE),
        ('MMSE + ViT receiver', 'Undoes the fading of each symbol using channel state, then rebuilds the '
                                'latent.', BLUE),
    ]):
        card(s, MARGIN + i * 4.1, 4.6, 3.85, 2.15, h, b, heading_color=col, size=16)

    # 7. Channel and receiver model -------------------------------------------
    s = d.content('System design', 'Channel and receiver model',
                  'Each symbol is multiplied by a fading coefficient h and corrupted by complex Gaussian noise. '
                  'Average symbol power is one, so the SNR is one over the noise variance. AWGN has h equal to '
                  'one; Rayleigh draws h from a zero-mean complex Gaussian (random phase and gain, no direct '
                  'path); Rician adds a strong line-of-sight component with K-factor 3. The receiver knows h '
                  'and the SNR and applies the MMSE equalizer before the transformer.')
    eqs = [
        ('Channel', [[('y = h · x + n,   n ~ CN(0, σ', {}), ('2', {'sup': True}), (')', {})]]),
        ('Power and SNR', [[('E|x|', {}), ('2', {'sup': True}), (' = 1,   SNR = 1 / σ', {}), ('2', {'sup': True})]]),
        ('MMSE equalizer', [[('x', {}), ('eq', {'sub': True}), (' = y · h* / ( |h|', {}), ('2', {'sup': True}),
                             (' + σ', {}), ('2', {'sup': True}), (' )', {})]]),
    ]
    for i, (h, eq) in enumerate(eqs):
        y = 1.8 + i * 1.65
        rect(s, MARGIN, y, 5.9, 1.45)
        text(s, MARGIN + 0.3, y + 0.22, 5.3, 0.35, h.upper(), size=12, color=BLUE, bold=True)
        text(s, MARGIN + 0.3, y + 0.62, 5.3, 0.6, eq, size=22, color=INK)
    chans = [('AWGN', 'h = 1. Noise only; the baseline case.', MUTED),
             ('Rician — line of sight', 'Strong direct path plus scattering (K-factor 3).', AQUA),
             ('Rayleigh — non line of sight', 'h ~ CN(0, 1): random phase and gain on every symbol.', ORANGE)]
    for i, (h, b, col) in enumerate(chans):
        y = 1.8 + i * 1.65
        rect(s, 6.75, y, 5.95, 1.45, fill=WHITE, line=CARD_EDGE)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.05), Inches(y + 0.3), Inches(0.24), Inches(0.24))
        dot.fill.solid(); dot.fill.fore_color.rgb = col; dot.line.fill.background(); dot.shadow.inherit = False
        text(s, 7.45, y + 0.2, 5.1, 0.42, h, size=19, bold=True)
        text(s, 7.45, y + 0.68, 5.1, 0.65, b, size=16, color=INK_2)
    text(s, MARGIN, 6.68, 12.1, 0.3, 'E|h|² = 1 for all fading models · h drawn independently per symbol · '
                                     'receiver knows h and the SNR', size=13, color=MUTED)

    # 8. Training -------------------------------------------------------------
    s = d.content('Methodology', 'Training methodology',
                  'The VAE is frozen, so we encode every training image once and cache the latents. Training '
                  'then only runs the two transformers. Each batch picks a random channel type, each sample a '
                  'random SNR between 0 and 20 dB, so one model learns all conditions. The loss is latent MSE '
                  'plus a small cosine term. Training images: CIFAR-10 (upsampled) and crops of DIV2K photos. '
                  'Testing uses the standard Kodak images.')
    steps = [('Cache latents', 'Encode training images once with the frozen VAE'),
             ('Random channel', 'Per batch: AWGN, Rician or Rayleigh. Per sample: SNR 0–20 dB'),
             ('Optimize', 'Latent MSE + 0.1 · (1 − cosine similarity), AdamW, 30 epochs'),
             ('Evaluate', 'Kodak test images, PSNR and SSIM per channel and SNR')]
    for i, (h, b) in enumerate(steps):
        x = MARGIN + i * 3.1
        rect(s, x, 1.8, 2.75, 2.75)
        circle_num(s, x + 0.25, 2.05, i + 1)
        text(s, x + 0.25, 2.75, 2.3, 0.45, h, size=19, bold=True)
        text(s, x + 0.25, 3.22, 2.3, 1.25, b, size=15, color=INK_2)
        if i < 3:
            arrow(s, x + 2.78, 3.0, w=0.29, h=0.34)
    stats = [('30', 'epochs per model'), ('1–3 h', 'per model on an RTX 3060 laptop GPU'),
             (f"~{round(N['fps']['k1024_div2k'])} fps", 'end-to-end inference')]
    for i, (big, small) in enumerate(stats):
        x = MARGIN + i * 4.1
        text(s, x, 4.95, 3.85, 0.85, big, size=40, color=BLUE, bold=True)
        text(s, x, 5.8, 3.85, 0.5, small, size=15, color=INK_2)
    text(s, MARGIN, 6.45, 12.1, 0.35, 'Training data: CIFAR-10 (upsampled to 128x128) and DIV2K photo crops   ·   '
                                      'Test data: Kodak (144 tiles of 128x128)', size=13, color=MUTED)

    # 9. Evolution ------------------------------------------------------------
    s = d.content('Evolution', 'How the design evolved',
                  'The first prototype had the full pipeline and a live demo, but a careful evaluation showed '
                  'it had not learned Rayleigh fading, its SNR labels were 3 dB off, and output was blurry. '
                  'We redesigned the receiver, fixed the SNR definition, made bandwidth configurable, trained '
                  'on real photographs, and built a proper evaluation with digital baselines.')
    cols = [('Initial prototype', NAVY_2, ['Full pipeline and live webcam demo', '256 channel symbols per frame',
                                           'Trained on CIFAR-10 only']),
            ('Issues found', ORANGE, ['Rayleigh not learned: error flat from 0 to 40 dB',
                                      'SNR labels 3 dB optimistic', 'Blurry output, no image-space evaluation']),
            ('Redesign', BLUE, ['MMSE equalization before the transformer', 'Exact SNR; configurable bandwidth',
                                'Real-photo training; Kodak evaluation vs digital'])]
    for i, (h, col, items_) in enumerate(cols):
        x = MARGIN + i * 4.2
        top = rect(s, x, 1.85, 3.7, 0.75, fill=col)
        text(s, x, 1.85, 3.7, 0.75, h, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        rect(s, x, 2.75, 3.7, 3.85)
        bullets(s, x + 0.28, 3.0, 3.2, 3.5, items_, size=17, gap=12)
        if i < 2:
            arrow(s, x + 3.78, 2.03, w=0.34, h=0.38)

    # 10. Result 1 ------------------------------------------------------------
    raw0, raw20 = psnr('k256_raw', 'rayleigh', 0), psnr('k256_raw', 'rayleigh', 20)
    mm0, mm20 = psnr('k256', 'rayleigh', 0), psnr('k256', 'rayleigh', 20)
    s = d.content('Result 1', 'Explicit equalization makes Rayleigh fading learnable',
                  f'Same model, data and training; only the receiver input differs. Feeding raw received symbols '
                  f'and channel coefficients into the transformer gives {raw0:.1f} dB at 0 dB SNR and only '
                  f'{raw20:.1f} dB at 20 dB: a hundredfold better signal buys about half a decibel. With MMSE '
                  f'equalization first, quality rises from {mm0:.1f} to {mm20:.1f} dB. The reason: undoing fading '
                  f'needs a per-symbol division by h, which the transformer\'s linear input layer cannot compute.')
    picture(s, os.path.join(FIG_SLIDE, 'fig2_equalization.png'), MARGIN, 1.7, 7.7, 4.9)
    x0 = 8.65
    text(s, x0, 1.9, 4.1, 0.85, f'{raw20:.1f} dB', size=44, color=ORANGE, bold=True)
    text(s, x0, 2.75, 4.1, 0.6, 'raw channel input at 20 dB SNR — almost flat across SNR', size=15,
         color=INK_2)
    text(s, x0, 3.55, 4.1, 0.85, f'{mm20:.1f} dB', size=44, color=BLUE, bold=True)
    text(s, x0, 4.4, 4.1, 0.6, f'MMSE-equalized at 20 dB, rising steadily from {mm0:.1f} dB', size=15,
         color=INK_2)
    rect(s, x0, 5.25, 4.05, 1.45, fill=RGBColor(0xE8, 0xF1, 0xFC))
    text(s, x0 + 0.2, 5.25, 3.65, 1.45, 'Undoing fading needs a per-symbol division (y / h) that a linear '
                                        'layer cannot compute', size=15, anchor=MSO_ANCHOR.MIDDLE)

    # 11. Result 2 ------------------------------------------------------------
    s = d.content('Result 2', 'More channel symbols per frame give sharper images',
                  'K is the number of complex channel symbols per frame. Doubling it twice, from 256 to 1024, '
                  'improves quality at every SNR, and the gain grows with SNR: about 1 dB at 0 dB and 3 dB at '
                  '20 dB. At low SNR, noise is the bottleneck; at high SNR, the number of symbols is. All curves '
                  'stay below the VAE ceiling of 26.9 dB.')
    picture(s, os.path.join(FIG_SLIDE, 'fig3_bandwidth.png'), MARGIN, 1.7, 7.4, 4.9)
    rows = [('K', 'Share of pixels', '0 dB', '10 dB', '20 dB')]
    for k, ratio in [(256, '1/192'), (512, '1/96'), (1024, '1/48')]:
        tag = 'k256' if k == 256 else f'k{k}'
        rows.append((str(k), ratio, *(f"{psnr(tag, 'rayleigh', s_):.1f}" for s_ in (0, 10, 20))))
    rows.append(('VAE only', '—', *([f"{N['vae_ceiling_psnr']:.1f}"] * 3)))
    tbl = s.shapes.add_table(len(rows), 5, Inches(8.3), Inches(1.95), Inches(4.4), Inches(2.6)).table
    for c, wdt in enumerate([0.95, 1.05, 0.8, 0.8, 0.8]):
        tbl.columns[c].width = Inches(wdt)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (CARD if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = val
            run.font.name, run.font.size = FONT, Pt(14)
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else INK
    text(s, 8.3, 4.62, 4.4, 0.3, 'PSNR in dB, Rayleigh fading, Kodak', size=12, color=MUTED)
    g0 = psnr('k1024', 'rayleigh', 0) - psnr('k256', 'rayleigh', 0)
    g20 = psnr('k1024', 'rayleigh', 20) - psnr('k256', 'rayleigh', 20)
    bullets(s, 8.3, 5.15, 4.4, 1.6, [f'256 -> 1024 symbols: +{g0:.1f} dB at 0 dB, +{g20:.1f} dB at 20 dB',
                                      'Low SNR: noise limits quality. High SNR: bandwidth does'], size=15)

    # 12. Result 3 ------------------------------------------------------------
    j0 = [psnr('k1024_div2k', ch, 0) for ch in ('awgn', 'rician', 'rayleigh')]
    s = d.content('Result 3', 'Graceful degradation versus the digital cliff',
                  'Every scheme uses the same bandwidth, 1024 symbols per frame. A digital link with its rate '
                  'fixed for 10 dB fails completely below 10 dB and does not improve above it. Even with rate '
                  f'adaptation, a practical digital link fails at 0 dB, where Deep JSCC still delivers '
                  f'{min(j0):.0f} to {max(j0):.0f} dB. From about 2 to 4 dB upward, depending on the channel, the adaptive '
                  'digital link is better, by 2.5 to 3.3 dB at 20 dB, because our quality is capped by the VAE at 26.9 dB. We state this honestly.')
    picture(s, os.path.join(FIG_SLIDE, 'fig4_vs_digital.png'), MARGIN, 1.6, 12.1, 3.95)
    takeaways = [('Below the design SNR', 'Fixed-rate digital fails (grey frame); Deep JSCC keeps working', BLUE),
                 ('At 0 dB', f'All digital links fail; Deep JSCC still gives {min(j0):.0f}–{max(j0):.0f} dB',
                  BLUE),
                 ('Above ~2–4 dB', 'Adaptive digital wins (up to ~3 dB); we are capped by the VAE (26.9 dB)',
                  ORANGE)]
    for i, (h, b, col) in enumerate(takeaways):
        x = MARGIN + i * 4.1
        rect(s, x, 5.65, 3.85, 1.22)
        text(s, x + 0.22, 5.75, 3.45, 0.38, h, size=16, color=col, bold=True)
        text(s, x + 0.22, 6.12, 3.45, 0.7, b, size=14, color=INK_2)

    # 13. Visual results ------------------------------------------------------
    s = d.content('Result 4', 'What the reconstructions look like',
                  'Rayleigh fading, 1024 symbols. At 20 dB both look good; digital is slightly sharper. At 10 dB, '
                  'Deep JSCC keeps fine detail such as the face paint and feathers while digital becomes blocky, '
                  'although PSNR slightly favours digital there. At 0 dB, Deep JSCC is degraded but recognisable '
                  'and digital fails completely. The bottom row is a real webcam frame from our live system.')
    pic = picture(s, os.path.join(FIG_SLIDE, 'fig5_visual.png'), MARGIN, 2.25, 8.8, 4.65, align='left')
    # Column labels as slide text (the image's own labels are too small when projected).
    # Grid geometry from make_figures.fig_visual: 7 columns, left/right 0.01, wspace 0.04.
    pw = pic.width / 914400
    col_w = 0.98 / (7 + 6 * 0.04)
    labels = [['Original'], ['JSCC', '20 dB'], ['JSCC', '10 dB'], ['JSCC', '0 dB'],
              ['Digital', '20 dB'], ['Digital', '10 dB'], ['Digital', '0 dB']]
    for i, lab in enumerate(labels):
        cx = MARGIN + pw * (0.01 + i * col_w * 1.04 + col_w / 2)
        color = INK if i == 0 else (BLUE if i <= 3 else ORANGE)
        text(s, cx - 0.62, 1.6, 1.24, 0.62, lab, size=13, color=color, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, line_spacing=0.95)
    obs = [('20 dB', 'Both clear; digital slightly sharper'),
           ('10 dB', 'Deep JSCC keeps fine detail; digital turns blocky'),
           ('0 dB', 'Deep JSCC degraded but recognisable; digital fails')]
    for i, (h, b) in enumerate(obs):
        y = 1.75 + i * 1.62
        rect(s, 9.65, y, 3.05, 1.42)
        text(s, 9.87, y + 0.18, 2.6, 0.42, h, size=20, color=BLUE, bold=True)
        text(s, 9.87, y + 0.62, 2.65, 0.75, b, size=14, color=INK_2)
    text(s, 9.65, 6.6, 3.05, 0.3, 'Rayleigh, K = 1024; bottom row: webcam', size=11, color=MUTED)

    # 14. Conclusion ----------------------------------------------------------
    s = d.slide('Summarize: one model across three channels, the equalization insight, a real-time system with '
                'an honest comparison against digital. Limitations: simulated channel with perfect channel '
                'knowledge, 128x128 frames, quality capped by the VAE. Future work: generative decoders for '
                'sharper low-SNR output, temporal conditioning for video, and over-the-air tests.', dark=True)
    text(s, MARGIN, 0.38, 12.1, 0.3, 'CONCLUSION', size=12, color=ICE, bold=True)
    text(s, MARGIN, 0.68, 12.1, 0.8, 'Conclusion and future work', size=32, color=WHITE, bold=True)
    gain = psnr('k256', 'rayleigh', 20) - psnr('k256_raw', 'rayleigh', 20)
    blocks = [('Contributions', ['One ViT JSCC model for AWGN, Rician and Rayleigh',
                                 f'Equalizing before the transformer: +{gain:.1f} dB on Rayleigh at 20 dB',
                                 f"Real-time system (~{round(N['fps']['k1024_div2k'])} fps) with digital baselines"]),
              ('Limitations', ['Simulated channel, perfect channel knowledge', '128x128 frames; VAE caps quality '
                               f"at {N['vae_ceiling_psnr']:.1f} dB", 'Adaptive digital wins at high SNR']),
              ('Future work', ['Generative decoder for sharp low-SNR images', 'Temporal conditioning for video',
                               'Over-the-air tests with software-defined radios'])]
    for i, (h, items_) in enumerate(blocks):
        x = MARGIN + i * 4.1
        rect(s, x, 1.8, 3.85, 4.25, fill=NAVY_2)
        text(s, x + 0.3, 2.05, 3.3, 0.5, h, size=22, color=WHITE, bold=True)
        bullets(s, x + 0.3, 2.75, 3.3, 3.9, items_, size=17, color=WHITE, gap=14)
    text(s, MARGIN, 7.02, 9, 0.3, f'{TEAM}  |  {TITLE}', size=10, color=ICE)
    text(s, 11.733, 7.02, 1, 0.3, str(d.count), size=10, color=ICE, align=PP_ALIGN.RIGHT)

    # 15. References and thank you --------------------------------------------
    s = d.slide('Key references for the methods we build on and compare against. Thank the panel.', dark=True)
    text(s, MARGIN, 0.38, 6.6, 0.3, 'REFERENCES', size=12, color=ICE, bold=True)
    refs = [
        'E. Bourtsoulatze, D. B. Kurka, D. Gündüz, "Deep joint source-channel coding for wireless image '
        'transmission," IEEE Trans. Cogn. Commun. Netw., 2019.',
        'J. Xu, B. Ai, W. Chen, A. Yang, P. Sun, M. Rodrigues, "Wireless image transmission using deep source '
        'channel coding with attention modules," IEEE Trans. Circuits Syst. Video Technol., 2022.',
        'K. Yang, S. Wang, J. Dai, K. Tan, K. Niu, P. Zhang, "WITT: A wireless image transmission transformer '
        'for semantic communications," ICASSP, 2023.',
        'R. Rombach et al., "High-resolution image synthesis with latent diffusion models," CVPR, 2022.',
        'A. Dosovitskiy et al., "An image is worth 16x16 words: Transformers for image recognition at scale," '
        'ICLR, 2021.',
        'C. E. Shannon, "A mathematical theory of communication," Bell Syst. Tech. J., 1948.',
    ]
    text(s, MARGIN, 0.9, 6.9, 6.0, [[(f'[{i + 1}]  ', {'bold': True, 'color': ICE}), (r, {})]
                                     for i, r in enumerate(refs)], size=13, color=WHITE, line_spacing=1.15)
    for p in s.shapes[-1].text_frame.paragraphs:
        p.space_after = Pt(10)
    text(s, 8.0, 2.55, 4.8, 1.4, 'Thank you', size=60, color=WHITE, bold=True)
    text(s, 8.0, 3.95, 4.8, 0.5, TEAM, size=22, color=ICE, bold=True)
    text(s, 8.0, 4.5, 4.8, 1.6, MEMBERS, size=16, color=ICE, line_spacing=1.2)

    d.save()
    print(f'Wrote {OUT} ({d.count} slides)')


if __name__ == '__main__':
    build()
